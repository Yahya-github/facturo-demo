import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.units import pixels_to_EMU

from facturo import paths
from facturo.brand import TPS_NUMBER, TVQ_NUMBER
from facturo.invoicing import discounts

PDF_TIMEOUT = 120

# Custom Excel table style "Facture avec frais financiers (en bleu)" uses
# theme accent2 (#58B6C0). LibreOffice ignores custom table styles when it
# exports to PDF, so we bake the blue look directly into the cells:
#   - banded rows   : accent2 tint +0.8  -> light blue fill
#   - header text   : accent2 tint -0.5  -> dark teal
#   - table borders : accent2            -> blue lines
TABLE_BLUE = "FF58B6C0"
TABLE_BLUE_DARK = "FF2C5B60"
TABLE_STRIPE_FILL = PatternFill("solid", fgColor="FFDEF0F2")
TABLE_BORDER_BLUE = Side(style="thin", color=TABLE_BLUE)

# Template ships inside the program (read-only); generated invoices live next
# to the exe so the client can find them.
TEMPLATE_PATH = paths.template_path()
OUTPUT_DIR = paths.output_dir()

TABLE_NAME = "Données"
SHEET_NAME = "FRAIS FINANCIERS"
HEADER_ROW = 12
FIRST_DATA_ROW = 13
TEMPLATE_LAST_DATA_ROW = 38
TEMPLATE_DATA_COUNT = TEMPLATE_LAST_DATA_ROW - FIRST_DATA_ROW + 1

CALIBRI_28 = Font(name="Calibri", size=28)
CALIBRI_28_BOLD = Font(name="Calibri", size=28, bold=True)
CALIBRI_28_BOLD_ITALIC = Font(name="Calibri", size=28, bold=True, italic=True)
CALIBRI_36_BOLD_ITALIC = Font(name="Calibri", size=36, bold=True, italic=True)
CENTURY_28 = Font(name="Century Gothic", size=28)
CENTURY_28_BOLD_ITALIC = Font(name="Century Gothic", size=28, bold=True, italic=True)
CENTURY_18_BOLD_ITALIC = Font(name="Century Gothic", size=18, bold=True, italic=True)

FMT_QTY = '#,##0.00'
FMT_RATE = '"$"#,##0.00'
FMT_TOTAL = '0.00 "$"'
FMT_DATE = 'dd-mm-yyyy'

# Company logo placement. The logo fills the large empty block on the right of
# the header — column E, from just under the "RÉF CLIENT" line down to the gap
# above the "Données" table. The image is fit inside that block (aspect ratio
# preserved) and centered horizontally, so it reads as a prominent letterhead
# logo. LibreOffice renders the embedded image into the PDF as well.
LOGO_BOX_MAX_WIDTH_PX = 660    # ~ usable width of column E
LOGO_BOX_MAX_HEIGHT_PX = 330   # tall enough to fill the block, with a gap above the table
LOGO_ANCHOR_COL = 4            # 0-based -> column E
LOGO_ANCHOR_ROW = 5            # 0-based -> row 6, just below the "RÉF CLIENT" line
LOGO_TOP_MARGIN_PX = 6
# Approx pixel width of column E (openpyxl width 98.1 -> ~ width*7 + 12 px).
LOGO_ANCHOR_COL_WIDTH_PX = 698


def _fmt_money(n: float) -> str:
    """Mirror the in-app money format: space thousands, period decimals, ' $'."""
    return f"{n:,.2f}".replace(",", " ") + " $"


def _fmt_percent(valeur) -> str:
    v = float(valeur)
    s = f"{v:.2f}".rstrip("0").rstrip(".")
    return f"{s} %"


def generate_invoice(
    client: dict,
    numero: str,
    date_str: str,
    billets: list[dict],
    output_filename: str | None = None,
    remise_pct: float = 0.0,
    remise_montant: float = 0.0,
) -> Path:
    """Write an invoice xlsx. Discount amounts all come from ``discounts``:
    billet rows carry their own discount only; the invoice-level percent and
    fixed amount each get their own REMISE line under SOUS-TOTAL."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(str(TEMPLATE_PATH))
    ws = wb[SHEET_NAME]

    ws["E3"] = numero
    ws["E3"].font = CENTURY_28

    invoice_date = datetime.strptime(date_str, "%Y-%m-%d")
    ws["E4"] = invoice_date
    ws["E4"].number_format = FMT_DATE
    ws["E4"].font = CENTURY_28

    ws["E5"] = client["ref"]
    ws["E5"].font = CENTURY_28_BOLD_ITALIC

    ws["D8"] = client["nom"]
    ws["D8"].font = CENTURY_18_BOLD_ITALIC

    ws["B9"] = client["nom"]
    ws["B9"].font = CENTURY_28_BOLD_ITALIC

    ws["B10"] = client["adresse"]
    ws["B10"].font = Font(name="Century Gothic", size=28, italic=True)

    num_billets = len(billets)

    _clear_template_data(ws)
    _resize_table(ws, num_billets)
    _unhide_data_rows(ws, num_billets)
    _fill_billet_rows(ws, billets)
    _apply_blue_table_style(ws, num_billets)
    facture = {"remise_pct": remise_pct, "remise_montant": remise_montant}
    total_row = _write_totals_block(ws, num_billets, billets, facture)
    _set_print_layout(ws, total_row)
    _insert_logo(ws)

    if output_filename is None:
        output_filename = invoice_filename(client, numero, date_str)

    out_path = OUTPUT_DIR / output_filename
    wb.save(str(out_path))

    return out_path


def invoice_filename(client: dict, numero: str, date_str: str) -> str:
    """Default xlsx name of an invoice. Public so a caller can see which file a
    generation is about to overwrite (and put it back if the save fails)."""
    safe_nom = "".join(c if c.isalnum() or c in "-_ " else "" for c in client["nom"])
    safe_nom = safe_nom.strip().replace(" ", "_")[:30]
    return f"facture_{numero}_{safe_nom}_{date_str}.xlsx"


def _clear_template_data(ws):
    for row_idx in range(FIRST_DATA_ROW, TEMPLATE_LAST_DATA_ROW + 1):
        for col in range(2, 7):
            ws.cell(row=row_idx, column=col).value = None

    # Clear two rows past the four template total rows: the invoice-level
    # percent and fixed-amount discounts each insert a REMISE line, pushing
    # TOTAL DÛ down by up to two.
    for row_idx in range(TEMPLATE_LAST_DATA_ROW + 1, TEMPLATE_LAST_DATA_ROW + 7):
        for col in range(2, 7):
            ws.cell(row=row_idx, column=col).value = None
            ws.cell(row=row_idx, column=col).font = CALIBRI_28
            ws.cell(row=row_idx, column=col).number_format = "General"


def _resize_table(ws, num_billets: int):
    last_data_row = FIRST_DATA_ROW + num_billets - 1

    if num_billets > TEMPLATE_DATA_COUNT:
        extra = num_billets - TEMPLATE_DATA_COUNT
        ws.insert_rows(TEMPLATE_LAST_DATA_ROW + 1, extra)
    elif num_billets < TEMPLATE_DATA_COUNT:
        remove = TEMPLATE_DATA_COUNT - num_billets
        ws.delete_rows(FIRST_DATA_ROW + num_billets, remove)

    new_ref = f"B{HEADER_ROW}:F{last_data_row}"
    for table in ws.tables.values():
        if table.name == TABLE_NAME:
            table.ref = new_ref
            # The autofilter range must stay equal to the table range, or
            # Excel (and especially Excel Online) rejects the file as corrupt.
            if table.autoFilter is not None:
                table.autoFilter.ref = new_ref
            break

    return last_data_row


def _unhide_data_rows(ws, num_billets: int):
    for i in range(num_billets):
        row_idx = FIRST_DATA_ROW + i
        ws.row_dimensions[row_idx].hidden = False


def _billet_discount_label(billet: dict, gross: float, net: float) -> str:
    """Explain why TOTAL is below qty × rate, e.g. 'Remise 10 % + 50.00 $ : -150.00 $'."""
    pct, montant = discounts.normalize_discount(billet)
    parts = []
    if pct > 0:
        parts.append(_fmt_percent(pct))
    if montant > 0:
        parts.append(_fmt_money(montant))
    return f"Remise {' + '.join(parts)} : -{_fmt_money(gross - net)}"


def _fill_billet_rows(ws, billets: list[dict]):
    for i, billet in enumerate(billets):
        row_idx = FIRST_DATA_ROW + i

        qty = float(billet["quantite"])
        rate = float(billet["taux"])
        gross = discounts.billet_gross(billet)
        net = discounts.billet_net(billet)

        # Two ways to describe a line: the structured billet fields (default),
        # or a free-text description the user typed (e.g. "Réparation").
        description_parts = []
        if billet.get("desc_libre"):
            if billet.get("description"):
                description_parts.append(str(billet["description"]))
        else:
            if billet.get("chantier"):
                description_parts.append(f"Client: {billet['chantier']}")
            if billet.get("plaque"):
                description_parts.append(f"Plaque: {billet['plaque']}")
            if billet.get("date_billet"):
                description_parts.append(f"Date: {billet['date_billet']}")
            if billet.get("numero_billet"):
                description_parts.append(f"# Billet: {billet['numero_billet']}")
        # Spell out a per-billet discount on the line itself, so the customer can
        # see why TOTAL is below quantité × montant.
        if net < gross:
            description_parts.append(_billet_discount_label(billet, gross, net))

        cell_b = ws.cell(row=row_idx, column=2)
        cell_b.value = "\n".join(description_parts)
        cell_b.font = CALIBRI_28
        cell_b.alignment = Alignment(wrap_text=True, vertical="center")

        cell_c = ws.cell(row=row_idx, column=3)
        cell_c.value = qty
        cell_c.font = CALIBRI_28
        cell_c.number_format = FMT_QTY
        cell_c.alignment = Alignment(horizontal="center", vertical="center")

        cell_d = ws.cell(row=row_idx, column=4)
        cell_d.value = rate
        cell_d.font = CALIBRI_28
        cell_d.number_format = FMT_RATE
        cell_d.alignment = Alignment(horizontal="center", vertical="center")

        cell_e = ws.cell(row=row_idx, column=5)
        cell_e.value = net
        cell_e.font = CALIBRI_28
        cell_e.number_format = FMT_TOTAL
        cell_e.alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=row_idx, column=6).value = None


def _apply_blue_table_style(ws, num_billets: int):
    """Bake the blue Excel table look into cells so it survives PDF export.

    Header row gets dark-teal bold text and a blue underline; data rows get
    alternating light-blue banding; the table is framed top and bottom in blue.
    """
    last_data_row = FIRST_DATA_ROW + num_billets - 1

    # Header row: keep the existing 28pt size, recolor + bold, blue underline.
    # Vertical separators on every cell split the categories apart.
    for col in range(2, 7):
        cell = ws.cell(row=HEADER_ROW, column=col)
        base = cell.font
        cell.font = Font(name=base.name, size=base.size, bold=True, color=TABLE_BLUE_DARK)
        cell.border = Border(
            left=TABLE_BORDER_BLUE,
            right=TABLE_BORDER_BLUE,
            top=TABLE_BORDER_BLUE,
            bottom=TABLE_BORDER_BLUE,
        )
        # Center the numeric column titles over their values (QUANTITÉ,
        # MONTANT, TOTAL). The template padded TOTAL with leading spaces to
        # fake centering — strip that so real centering lines up.
        if col in (3, 4, 5):
            if isinstance(cell.value, str):
                cell.value = cell.value.strip()
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Excel requires each table column's name to match its header cell text
    # exactly; after stripping the padded TOTAL title above, mirror the cell
    # values back into the table definition or the file won't open.
    for table in ws.tables.values():
        if table.name == TABLE_NAME:
            for idx, table_col in enumerate(table.tableColumns):
                value = ws.cell(row=HEADER_ROW, column=2 + idx).value
                if isinstance(value, str):
                    table_col.name = value
            break

    # Data rows: band every other row, and frame every cell with vertical
    # separators so each category column is visually divided.
    for i in range(num_billets):
        row_idx = FIRST_DATA_ROW + i
        is_last = row_idx == last_data_row
        for col in range(2, 7):
            cell = ws.cell(row=row_idx, column=col)
            if i % 2 == 0:
                cell.fill = TABLE_STRIPE_FILL
            cell.border = Border(
                left=TABLE_BORDER_BLUE,
                right=TABLE_BORDER_BLUE,
                bottom=TABLE_BORDER_BLUE if is_last else None,
            )


def _write_totals_block(ws, num_billets: int, billets: list[dict], facture: dict) -> int:
    """Write SOUS-TOTAL / (REMISE x %) / (REMISE) / TVQ / TPS / TOTAL DÛ.

    SOUS-TOTAL is the sum of the billet rows printed above it; each REMISE
    line appears only when it removes something, and SOUS-TOTAL minus the
    REMISE lines is exactly the taxable base. Returns the TOTAL DÛ row.
    """
    last_data_row = FIRST_DATA_ROW + num_billets - 1
    totals = discounts.invoice_totals(billets, facture)
    pct, _montant = discounts.normalize_discount(facture)

    row = last_data_row + 1

    ws.cell(row=row, column=4).value = "SOUS-TOTAL:"
    ws.cell(row=row, column=4).font = CALIBRI_28_BOLD
    ws.cell(row=row, column=5).value = totals.sous_total
    ws.cell(row=row, column=5).font = CALIBRI_28_BOLD
    ws.cell(row=row, column=5).number_format = FMT_TOTAL
    row += 1

    remise_lines = (
        (f"REMISE ({_fmt_percent(pct)}):", totals.remise_pct_amount),
        ("REMISE:", totals.remise_montant_amount),
    )
    for label, amount in remise_lines:
        if amount <= 0:
            continue
        ws.cell(row=row, column=4).value = label
        ws.cell(row=row, column=4).font = CALIBRI_28_BOLD
        ws.cell(row=row, column=5).value = -amount
        ws.cell(row=row, column=5).font = CALIBRI_28_BOLD
        ws.cell(row=row, column=5).number_format = FMT_TOTAL
        row += 1

    ws.cell(row=row, column=2).value = f"TVQ: {TVQ_NUMBER}"
    ws.cell(row=row, column=2).font = CALIBRI_28_BOLD
    ws.cell(row=row, column=4).value = "TVQ"
    ws.cell(row=row, column=4).font = CALIBRI_28_BOLD
    ws.cell(row=row, column=5).value = totals.tvq
    ws.cell(row=row, column=5).font = CALIBRI_28_BOLD
    ws.cell(row=row, column=5).number_format = FMT_TOTAL
    row += 1

    ws.cell(row=row, column=2).value = f"TPS: {TPS_NUMBER}"
    ws.cell(row=row, column=2).font = CALIBRI_28_BOLD_ITALIC
    ws.cell(row=row, column=4).value = "TPS"
    ws.cell(row=row, column=4).font = CALIBRI_28_BOLD
    ws.cell(row=row, column=5).value = totals.tps
    ws.cell(row=row, column=5).font = CALIBRI_28_BOLD
    ws.cell(row=row, column=5).number_format = FMT_TOTAL
    row += 1

    ws.cell(row=row, column=4).value = "TOTAL DÛ:"
    ws.cell(row=row, column=4).font = CALIBRI_28_BOLD
    ws.cell(row=row, column=5).value = totals.total
    ws.cell(row=row, column=5).font = CALIBRI_36_BOLD_ITALIC
    ws.cell(row=row, column=5).number_format = FMT_TOTAL

    return row


def _set_print_layout(ws, total_row: int):
    """Normalize the print settings so the PDF export matches the live data.

    The template ships with two stale, hard-coded print settings that only ever
    surface on a multi-page invoice:

      * ``print_title_rows`` was ``$15:$15`` — the *third billet row*. LibreOffice
        treats "rows to repeat" as a header to reprint at the top of every page,
        so that one billet was stamped a second time at the top of page 2 (the
        phantom duplicate line), which also broke the blue/white banding there.
        Repeat the real table-header row instead, so page 2 gets column titles
        rather than a duplicated billet.
      * ``print_area`` was frozen at ``$A$1:$F$42``; an invoice long enough to
        push its totals past row 42 would have them clipped out of the PDF. Track
        the actual content extent instead (``total_row`` is the TOTAL DÛ line,
        which moves down by one per invoice-level REMISE line shown).
    """
    ws.print_title_rows = f"${HEADER_ROW}:${HEADER_ROW}"
    ws.print_area = f"$A$1:$F${total_row}"


# Marker labels that sit in column D right below the billet rows, used to
# detect where the data table ends when reading an invoice back.
_TOTALS_LABELS = {"tvq", "tps", "sous-total:", "total dû:"}


def _parse_description(text: str) -> dict:
    """Reconstruct billet fields from a description cell, any shipped format."""
    fields = {"chantier": "", "plaque": "", "date_billet": "", "numero_billet": ""}
    for raw in (text or "").split("\n"):
        line = raw.strip()
        if line.startswith("Client: "):
            fields["chantier"] = line[len("Client: "):].strip()
        elif line.startswith("Plaque: "):
            fields["plaque"] = line[len("Plaque: "):].strip()
        elif line.startswith("Date: "):
            fields["date_billet"] = line[len("Date: "):].strip()
        elif line.startswith("# Billet: "):
            fields["numero_billet"] = line[len("# Billet: "):].strip()
        elif line.startswith("Billet date: "):
            rest = line[len("Billet date: "):]
            date_part, _, num_part = rest.rpartition(": ")
            fields["date_billet"] = date_part.strip()
            fields["numero_billet"] = num_part.strip()
        elif line.startswith("Billet: "):
            fields["numero_billet"] = line[len("Billet: "):].strip()
    return fields


def read_billets_from_xlsx(xlsx_path: Path) -> list[dict]:
    """Recover the billet line items from a generated invoice xlsx.

    Used to edit older invoices whose billets were never stored in the DB. The
    description cell is parsed back into fields; quantity and rate come straight
    from their columns. Returns an empty list if nothing is recoverable.
    """
    wb = openpyxl.load_workbook(str(xlsx_path))
    ws = wb[SHEET_NAME]

    billets: list[dict] = []
    row = FIRST_DATA_ROW
    while True:
        desc = ws.cell(row=row, column=2).value
        qty = ws.cell(row=row, column=3).value
        rate = ws.cell(row=row, column=4).value

        if isinstance(rate, str) and rate.strip().lower() in _TOTALS_LABELS:
            break
        if qty is None and (desc is None or not str(desc).strip()):
            break

        fields = _parse_description(str(desc) if desc else "")
        fields["quantite"] = float(qty or 0)
        fields["taux"] = float(rate or 0)
        billets.append(fields)
        row += 1

    return billets


def _insert_logo(ws):
    """Anchor the company logo (if set) in the right-hand header block.

    No-op when no logo has been uploaded. The image keeps its aspect ratio, is
    fit inside the large empty block on the right of the header (column E, below
    the "RÉF CLIENT" line and above the table), and centered horizontally so it
    reads as a letterhead logo. LibreOffice renders the embedded image into the
    PDF, so the logo appears in both exports.
    """
    logo = paths.logo_path()
    if not logo.exists():
        return

    from PIL import Image as PILImage

    with PILImage.open(logo) as im:
        src_w, src_h = im.size
    if not src_w or not src_h:
        return

    # Fit the image inside the target block, preserving the aspect ratio.
    scale = min(LOGO_BOX_MAX_WIDTH_PX / src_w, LOGO_BOX_MAX_HEIGHT_PX / src_h)
    width_px = round(src_w * scale)
    height_px = round(src_h * scale)

    # Center the image horizontally within column E.
    col_off_px = max(0, (LOGO_ANCHOR_COL_WIDTH_PX - width_px) // 2)

    marker = AnchorMarker(
        col=LOGO_ANCHOR_COL,
        colOff=pixels_to_EMU(col_off_px),
        row=LOGO_ANCHOR_ROW,
        rowOff=pixels_to_EMU(LOGO_TOP_MARGIN_PX),
    )
    size = XDRPositiveSize2D(pixels_to_EMU(width_px), pixels_to_EMU(height_px))

    # The generic "FACTURE" masthead label lives in the top-right corner
    # (merged D1:E1). Blank it so the header stays clean above the logo.
    ws["D1"] = None

    img = XLImage(str(logo))
    img.anchor = OneCellAnchor(_from=marker, ext=size)
    ws.add_image(img)


# Windows installs LibreOffice here but does NOT add it to PATH, so shutil.which
# fails. Check the standard install locations explicitly.
_WINDOWS_SOFFICE_PATHS = (
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
)

# LibreOffice shipped alongside the exe so PDF export works on a brand-new PC
# with no LibreOffice installed. Relative to app_dir() (the folder holding
# Factures.exe). Several layouts are accepted so the packager can drop either
# LibreOfficePortable or a plain copied LibreOffice "program" folder.
_BUNDLED_SOFFICE_RELPATHS = (
    Path("LibreOfficePortable") / "App" / "libreoffice" / "program" / "soffice.exe",
    Path("LibreOffice") / "program" / "soffice.exe",
    Path("libreoffice") / "program" / "soffice.exe",
    # Same layouts without the .exe suffix, for non-Windows dev/testing.
    Path("LibreOfficePortable") / "App" / "libreoffice" / "program" / "soffice",
    Path("LibreOffice") / "program" / "soffice",
)


def _find_bundled_soffice() -> str | None:
    """Locate a LibreOffice copy shipped next to the exe, if present."""
    base = paths.app_dir()
    for rel in _BUNDLED_SOFFICE_RELPATHS:
        candidate = base / rel
        if candidate.exists():
            return str(candidate)
    return None


def find_soffice() -> str | None:
    """Locate the LibreOffice executable, or None if unavailable.

    Search order: the copy bundled next to the exe (so a fresh PC works with no
    install), then PATH, then the standard Windows install locations.
    """
    bundled = _find_bundled_soffice()
    if bundled:
        return bundled
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if exe:
        return exe
    for candidate in _WINDOWS_SOFFICE_PATHS:
        if Path(candidate).exists():
            return candidate
    return None


def _find_soffice() -> str:
    exe = find_soffice()
    if not exe:
        raise RuntimeError(
            "LibreOffice est requis pour l'export PDF mais est introuvable. "
            "Installez LibreOffice (gratuit) depuis libreoffice.org, puis "
            "réessayez. L'export Excel fonctionne sans LibreOffice."
        )
    return exe


def convert_to_pdf(xlsx_path: Path, force: bool = False) -> Path:
    """Render an xlsx invoice to PDF with LibreOffice, preserving exact layout.

    The PDF is cached next to the xlsx and reused while it stays newer than
    the source file. Returns the path to the PDF.
    """
    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        raise FileNotFoundError(f"Fichier source introuvable: {xlsx_path}")

    pdf_path = xlsx_path.with_suffix(".pdf")
    if (
        not force
        and pdf_path.exists()
        and pdf_path.stat().st_mtime >= xlsx_path.stat().st_mtime
    ):
        return pdf_path

    soffice = _find_soffice()

    # Isolated profile dir so concurrent conversions don't clash on a shared
    # LibreOffice user profile lock.
    with tempfile.TemporaryDirectory() as tmp:
        profile_uri = (Path(tmp) / 'profile').as_uri()
        result = subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(xlsx_path.parent),
                str(xlsx_path),
                f"-env:UserInstallation={profile_uri}",
            ],
            capture_output=True,
            text=True,
            timeout=PDF_TIMEOUT,
        )

    if result.returncode != 0 or not pdf_path.exists():
        raise RuntimeError(
            f"Échec de la conversion PDF: {result.stderr.strip() or result.stdout.strip()}"
        )

    return pdf_path



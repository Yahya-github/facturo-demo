"""Integration tests for the discount fields across the HTTP surface:
POST /api/factures/generate, PUT /api/factures/{id}, GET /api/factures/{id},
and the xlsx file each of those writes to disk.

Every test runs against a throwaway database AND a throwaway output
directory (never the real ones): both facturo.core.database.DB_PATH and
facturo.invoicing.excel_generator.OUTPUT_DIR are module-level paths resolved
once at import time, so pointing them at a tmp_path before the app boots is
what isolates this from the user's real invoices (same pattern as
tests/test_known_values.py's `monkeypatch.setattr(db, "DB_PATH", ...)`).
"""

import json

import openpyxl
import pytest
from fastapi.testclient import TestClient

from facturo.core import database as db
from facturo.invoicing import excel_generator as gen
from facturo.server import app

EPSILON = 0.01  # xlsx values are money-rounded by nature; a cent is plenty


@pytest.fixture
def api_client(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(gen, "OUTPUT_DIR", tmp_path / "output")
    db.init_db()
    with TestClient(app) as c:
        yield c


@pytest.fixture
def client_id(api_client):
    resp = api_client.post(
        "/api/clients",
        json={"ref": "T", "prefix": "tst", "nom": "Client API Test", "adresse": "1 rue Test"},
    )
    assert resp.status_code == 200
    return resp.json()["id"]


def _facture_id_by_numero(api_client, client_id, numero):
    listing = api_client.get("/api/factures", params={"client_id": client_id}).json()
    return next(f["id"] for f in listing if f["numero"] == numero)


def _open_generated_workbook(filename):
    wb = openpyxl.load_workbook(str(gen.OUTPUT_DIR / filename))
    return wb[gen.SHEET_NAME]


def _read_totals_block(ws, num_billets):
    """Scan the rows right after the billet table for label -> value pairs,
    from SOUS-TOTAL down to and including TOTAL DÛ. Scans by label rather
    than assuming a fixed row offset, since the number of REMISE lines (0, 1
    or 2) shifts every row below it."""
    labels = {}
    row = gen.FIRST_DATA_ROW + num_billets
    for _ in range(20):
        label = ws.cell(row=row, column=4).value
        if label:
            labels[str(label).strip()] = ws.cell(row=row, column=5).value
            if str(label).strip().upper().startswith("TOTAL D"):
                break
        row += 1
    return labels


# ── Persistence round-trip ──────────────────────────────────────────────


def test_generate_persists_v2_fields_on_billet_and_invoice(api_client, client_id):
    resp = api_client.post(
        "/api/factures/generate",
        json={
            "client_id": client_id,
            "date": "2026-01-01",
            "billets": [
                {
                    "quantite": 10, "taux": 100, "remise_pct": 10, "remise_montant": 20,
                    "chantier": "Chantier A", "plaque": "L111111", "numero_billet": "1",
                },
                {
                    "quantite": 5, "taux": 80,
                    "chantier": "Chantier B", "plaque": "L222222", "numero_billet": "2",
                },
            ],
            "remise_pct": 5, "remise_montant": 15,
        },
    )
    assert resp.status_code == 200, resp.text
    numero = resp.json()["invoices"][0]["numero"]

    facture_id = _facture_id_by_numero(api_client, client_id, numero)
    body = api_client.get(f"/api/factures/{facture_id}").json()

    assert body["remise_pct"] == pytest.approx(5.0)
    assert body["remise_montant"] == pytest.approx(15.0)
    assert body["billets"][0]["remise_pct"] == pytest.approx(10.0)
    assert body["billets"][0]["remise_montant"] == pytest.approx(20.0)
    assert body["billets"][1]["remise_pct"] == pytest.approx(0.0)
    assert body["billets"][1]["remise_montant"] == pytest.approx(0.0)


def test_put_updates_v2_discount_fields(api_client, client_id):
    gen_resp = api_client.post(
        "/api/factures/generate",
        json={
            "client_id": client_id,
            "date": "2026-01-01",
            "billets": [
                {"quantite": 5, "taux": 100, "chantier": "A", "plaque": "L1", "numero_billet": "1"},
            ],
        },
    )
    assert gen_resp.status_code == 200, gen_resp.text
    numero = gen_resp.json()["invoices"][0]["numero"]
    facture_id = _facture_id_by_numero(api_client, client_id, numero)

    put_resp = api_client.put(
        f"/api/factures/{facture_id}",
        json={
            "client_id": client_id,
            "date": "2026-01-01",
            "billets": [
                {
                    "quantite": 5, "taux": 100, "remise_pct": 20, "remise_montant": 10,
                    "chantier": "A", "plaque": "L1", "numero_billet": "1",
                },
            ],
            "remise_pct": 5, "remise_montant": 3,
        },
    )
    assert put_resp.status_code == 200, put_resp.text

    body = api_client.get(f"/api/factures/{facture_id}").json()
    assert body["remise_pct"] == pytest.approx(5.0)
    assert body["remise_montant"] == pytest.approx(3.0)
    assert body["billets"][0]["remise_pct"] == pytest.approx(20.0)
    assert body["billets"][0]["remise_montant"] == pytest.approx(10.0)


def test_get_facture_normalizes_legacy_fields_for_the_editor(api_client, client_id):
    """A row inserted straight through the DB layer (bypassing any in-flight
    migration) with only legacy fields set must come back from GET with the
    equivalent v2 fields filled in, at both the invoice and billet level."""
    billets_json = json.dumps([
        {"quantite": 4, "taux": 50, "remise_type": "percent", "remise_valeur": 25},
    ])
    facture = db.save_facture(
        client_id, "leg001", "2026-01-01", "leg001.xlsx", billets_json=billets_json,
        remise_type="montant", remise_valeur=40.0,
    )

    body = api_client.get(f"/api/factures/{facture['id']}").json()

    assert body["remise_pct"] == pytest.approx(0.0)
    assert body["remise_montant"] == pytest.approx(40.0)
    assert body["billets"][0]["remise_pct"] == pytest.approx(25.0)
    assert body["billets"][0]["remise_montant"] == pytest.approx(0.0)


# ── xlsx output ──────────────────────────────────────────────────────────


def test_xlsx_totals_block_shows_separate_percent_and_montant_remise_lines(api_client, client_id):
    resp = api_client.post(
        "/api/factures/generate",
        json={
            "client_id": client_id,
            "date": "2026-01-01",
            "billets": [
                {
                    "quantite": 10, "taux": 100, "remise_pct": 10, "remise_montant": 20,
                    "chantier": "A", "plaque": "L1", "numero_billet": "1",
                },
                {
                    "quantite": 5, "taux": 80,
                    "chantier": "B", "plaque": "L2", "numero_billet": "2",
                },
            ],
            "remise_pct": 5, "remise_montant": 15,
        },
    )
    assert resp.status_code == 200, resp.text
    filename = resp.json()["invoices"][0]["filename"]

    ws = _open_generated_workbook(filename)
    totals = _read_totals_block(ws, num_billets=2)

    # billet1 net = 1000 - 100 - 20 = 880; billet2 net = 400 (no billet discount)
    assert totals["SOUS-TOTAL:"] == pytest.approx(1280.0, abs=EPSILON)

    pct_label = next((k for k in totals if k.upper().startswith("REMISE (")), None)
    assert pct_label is not None, f"no 'REMISE (x %):' line found among {list(totals)}"
    assert "5" in pct_label  # the invoice remise_pct value
    assert totals[pct_label] == pytest.approx(-64.0, abs=EPSILON)  # 1280 * 5%

    assert totals["REMISE:"] == pytest.approx(-15.0, abs=EPSILON)

    # base = 1280 - 64 - 15 = 1201; tps = 60.05; tvq = 119.79975
    assert totals["TOTAL DÛ:"] == pytest.approx(1380.85, abs=0.01)


def test_xlsx_omits_remise_lines_when_no_discount_is_set(api_client, client_id):
    resp = api_client.post(
        "/api/factures/generate",
        json={
            "client_id": client_id,
            "date": "2026-01-01",
            "billets": [
                {"quantite": 5, "taux": 100, "chantier": "A", "plaque": "L1", "numero_billet": "1"},
            ],
        },
    )
    assert resp.status_code == 200, resp.text
    filename = resp.json()["invoices"][0]["filename"]

    ws = _open_generated_workbook(filename)
    totals = _read_totals_block(ws, num_billets=1)

    assert not any(k.upper().startswith("REMISE") for k in totals)
    assert totals["SOUS-TOTAL:"] == pytest.approx(500.0, abs=EPSILON)


def test_xlsx_billet_row_sum_matches_printed_sous_total(api_client, client_id):
    """The SOUS-TOTAL cell must equal the sum of the individual billet row
    totals actually printed above it -- a customer must be able to verify it
    with a calculator. Invoice-level discounts show up ONLY in their own
    REMISE lines below, never baked back into a billet's own row."""
    resp = api_client.post(
        "/api/factures/generate",
        json={
            "client_id": client_id,
            "date": "2026-01-01",
            "billets": [
                {
                    "quantite": 10, "taux": 100, "remise_pct": 10,
                    "chantier": "A", "plaque": "L1", "numero_billet": "1",
                },
                {
                    "quantite": 5, "taux": 80, "remise_montant": 20,
                    "chantier": "B", "plaque": "L2", "numero_billet": "2",
                },
            ],
            "remise_pct": 5, "remise_montant": 100,
        },
    )
    assert resp.status_code == 200, resp.text
    filename = resp.json()["invoices"][0]["filename"]

    ws = _open_generated_workbook(filename)
    row1 = ws.cell(row=gen.FIRST_DATA_ROW, column=5).value
    row2 = ws.cell(row=gen.FIRST_DATA_ROW + 1, column=5).value

    assert row1 == pytest.approx(900.0, abs=EPSILON)  # 1000 - 10%
    assert row2 == pytest.approx(380.0, abs=EPSILON)  # 400 - 20

    totals = _read_totals_block(ws, num_billets=2)
    assert totals["SOUS-TOTAL:"] == pytest.approx(row1 + row2, abs=EPSILON)


def test_xlsx_billet_description_shows_both_discount_parts(api_client, client_id):
    resp = api_client.post(
        "/api/factures/generate",
        json={
            "client_id": client_id,
            "date": "2026-01-01",
            "billets": [
                {
                    "quantite": 10, "taux": 100, "remise_pct": 10, "remise_montant": 100,
                    "chantier": "A", "plaque": "L1", "numero_billet": "1",
                },
            ],
        },
    )
    assert resp.status_code == 200, resp.text
    filename = resp.json()["invoices"][0]["filename"]

    ws = _open_generated_workbook(filename)
    description = str(ws.cell(row=gen.FIRST_DATA_ROW, column=2).value or "")

    assert gen._fmt_percent(10) in description
    assert gen._fmt_money(100) in description
    # gross 1000, net 800 -> the actual total removed is 200
    assert f"-{gen._fmt_money(200)}" in description


def test_xlsx_billet_description_shows_montant_only_discount(api_client, client_id):
    resp = api_client.post(
        "/api/factures/generate",
        json={
            "client_id": client_id,
            "date": "2026-01-01",
            "billets": [
                {
                    "quantite": 5, "taux": 100, "remise_montant": 50,
                    "chantier": "A", "plaque": "L1", "numero_billet": "1",
                },
            ],
        },
    )
    assert resp.status_code == 200, resp.text
    filename = resp.json()["invoices"][0]["filename"]

    ws = _open_generated_workbook(filename)
    description = str(ws.cell(row=gen.FIRST_DATA_ROW, column=2).value or "")

    assert "Remise" in description
    assert "%" not in description.split("Remise", 1)[1].split("\n", 1)[0]
    assert gen._fmt_money(50) in description
